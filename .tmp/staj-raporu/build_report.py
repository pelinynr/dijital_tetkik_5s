from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

ROOT=Path(__file__).resolve().parents[2]
IMG=ROOT/'.tmp/staj-raporu/sunum'
OUT=ROOT/'Pelin_Yener_Staj_Raporu_Doldurulmus.docx'
d=Document(); sec=d.sections[0]
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2.5)
d.styles['Normal'].font.name='Times New Roman'; d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.line_spacing=1.5
for name,size in [('Title',20),('Heading 1',15),('Heading 2',12)]:
 d.styles[name].font.name='Times New Roman'; d.styles[name].font.size=Pt(size); d.styles[name].font.bold=True; d.styles[name].font.color.rgb=RGBColor(0,0,0)

def page(): d.add_page_break()
def tr_upper(t): return t.replace('i','İ').replace('ı','I').upper()
def heading(t,l=1): d.add_heading(tr_upper(t) if l==1 else t,level=l)
def para(t):
 p=d.add_paragraph(t); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; return p
def bullets(xs):
 for x in xs:d.add_paragraph(x,style='List Bullet')
def figure(no,slide,caption,w=15.2):
 p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(IMG/f'slide-{slide}.png'),width=Cm(w))
 p=d.add_paragraph(f'Şekil {no} {caption}');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1
 for r in p.runs:r.italic=True;r.font.size=Pt(9)
def table(headers,rows):
 t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for i,h in enumerate(headers):
  c=t.rows[0].cells[i];c.text=h;sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E1251B');c._tc.get_or_add_tcPr().append(sh)
  for r in c.paragraphs[0].runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row):cells[i].text=str(v)
 for row in t.rows:
  for c in row.cells:
   for p in c.paragraphs:p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(0)
 return t
def code(caption,text):
 p=d.add_paragraph(caption);p.runs[0].bold=True
 p=d.add_paragraph();p.paragraph_format.line_spacing=1;sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F1F3F5');p._p.get_or_add_pPr().append(sh)
 r=p.add_run(text);r.font.name='Consolas';r.font.size=Pt(8)

# Kapak
p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
logo=ROOT/'public/erdemir-logo.png'
if logo.exists():p.add_run().add_picture(str(logo),width=Cm(6))
for txt,size in [('DÜZCE ÜNİVERSİTESİ',16),('MÜHENDİSLİK FAKÜLTESİ',14),('BİLGİSAYAR MÜHENDİSLİĞİ BÖLÜMÜ',14),('',12),('STAJ RAPORU',20),('DİJİTAL 5S TETKİK YÖNETİM SİSTEMİ',17)]:
 p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(txt);r.bold=True;r.font.size=Pt(size)
table(['Öğrenci Bilgisi','Değer'],[['Öğrenci Numarası','231001053'],['Adı Soyadı','Pelin YENER'],['Bölümü','Bilgisayar Mühendisliği'],['Staj Kodu','BM499'],['Staj Tarihleri','10.08.2026 - 11.09.2026'],['Staj Birimi','ERP ve Çevresel Uygulamalar Müdürlüğü']])
p=d.add_paragraph('2026');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)

page();heading('İçindekiler')
for x in ['1. Giriş ve 5S Yaklaşımı','2. Kuruluş ve Çalışma Birimi','3. İhtiyaç Analizi','4. Projenin Tanımı ve Gereksinimler','5. Sistem Mimarisi','6. Frontend Geliştirme','7. Backend ve PostgreSQL Bağlantısı','8. Veritabanı Tasarımı','9. Kimlik Doğrulama ve Yetkilendirme','10. Kriter ve Planlama Modülü','11. QR Kodlu Tetkik ve Puanlama','12. Düzeltici Faaliyet Akışı','13. Raporlama ve Bildirim','14. Test ve Güvenlik','15. Sonuç','Ek A. Kod Blokları','Ek B. Panel Görselleri']:d.add_paragraph(x)

sections=[
('1. Giriş ve 5S Yaklaşımı',["Bu staj çalışmasında fabrika sahasındaki 5S tetkik sürecini dijital ortamda yönetmek amacıyla web tabanlı bir bilgi sistemi geliştirdim. Sistem; planlama, saha değerlendirmesi, fotoğraflı uygunsuzluk takibi, düzeltme, onay ve raporlama adımlarını tek uygulamada birleştirmektedir.","5S; çalışma alanının düzenli, temiz, güvenli ve sürdürülebilir biçimde yönetilmesini hedefleyen beş aşamalı bir yöntemdir. Ayıklama gereksiz malzemelerin uzaklaştırılmasını, Düzen kalan malzemelerin tanımlı yerlere konmasını, Temizlik alan ve ekipmanların temiz tutulmasını, Standartlaştırma ortak kuralları, Disiplin ise bu kuralların sürekliliğini ifade eder."],('1.1',2,'5S yönteminin temel adımları')),
('2. Kuruluş ve Çalışma Birimi',["Staj projesi Erdemir bünyesinde ERP ve Çevresel Uygulamalar Müdürlüğü ile ilişkili bir yazılım geliştirme konusu olarak yürütülmüştür. Çok sayıda üretim ünitesi ve sorumluluk rolü bulunduğu için 5S kayıtlarının merkezi olarak izlenmesi önemlidir.","Sistem; Çelikhane, Yüksek Fırın, Sıcak Haddehane, Soğuk Haddehaneler, Kok, Oksijen, Sinter ve Kireç Fabrikası gibi ünitelerin ayrı yetki sınırlarında yönetilmesini destekler. Proje doğrudan ERP entegrasyonu içermez; ancak gelecekte entegrasyonda kullanılabilecek düzenli veri ve REST servis altyapısı sağlar."],None),
('3. İhtiyaç Analizi',["Kriterler ve planlar Excel üzerinden hazırlanıyor, fotoğraflar ayrı ortamlarda saklanıyor ve uygunsuzluklar manuel yöntemlerle izleniyordu. Bu durum kayıt sürümlerinin karışmasına, görevlerin gecikmesine ve geçmiş analizlerin zorlaşmasına neden olabiliyordu.","Çözümde kayıtlar PostgreSQL üzerinde merkezileştirildi. Her kullanıcı yalnızca yetkili olduğu alan ve işlemleri görebilir. Önce ve sonra fotoğrafları aynı uygunsuzluk kaydıyla ilişkilendirilir."],('3.1',3,'Problem ve geliştirilen çözüm')),
('4. Projenin Tanımı ve Gereksinimler',["Dijital 5S Tetkik Sistemi, tetkik hazırlığından uygunsuzluğun kapanmasına kadar uçtan uca süreci yöneten rol tabanlı web uygulamasıdır.","Sistem dönemsel kriter oluşturma, görevli atama, admin onayı, alan QR kodu, Uygun/Uygun Değil değerlendirmesi, zorunlu fotoğraf, düzeltme onayı, bildirim ve geçmiş raporlama gereksinimlerini karşılamaktadır."],('4.1',4,'Projenin amacı')),
('5. Sistem Mimarisi',["Uygulama üç katmanlıdır. React arayüzü HTTP isteklerini Java Spring Boot REST servislerine gönderir. Backend iş kurallarını uygular ve PostgreSQL veritabanında kalıcı kayıt oluşturur."],('5.1',7,'Uygulamanın katmanlı mimarisi')),
('6. Frontend Geliştirme',["Frontend React ve TypeScript ile bileşen tabanlı geliştirilmiştir. HTML arayüz yapısını, CSS ise masaüstü ve mobil yerleşimi sağlar. Formlar React state yapısında tutulur ve Fetch API ile backend servislerine gönderilir.","Rol bilgisine göre sol menü değişir. Admin ve ünite sorumlusu yönetim panellerini, tetkikçi plan ve tetkik ekranlarını, alan sorumlusu ise düzeltme görevlerini görür."],('6.1',9,'Global admin genel bakış paneli')),
('7. Backend ve PostgreSQL Bağlantısı',["Backend Java ve Spring Boot ile geliştirilmiştir. Controller sınıfları REST isteklerini karşılar; PlanningService gibi servisler onay ve yetki kurallarını uygular; Spring JdbcTemplate parametrik SQL sorgularını çalıştırır.","PostgreSQL bağlantısı .env dosyasındaki DATABASE_URL değişkeninden okunur. Dijital5sApplication sınıfı PGSimpleDataSource nesnesini oluşturur. Kullanıcı adı ve parola kaynak kodun içine yazılmaz. Örnek biçim: jdbc:postgresql://localhost:5432/erdemir_5s?user=postgres&password=***"],None),
('8. Veritabanı Tasarımı',["users kullanıcı ve rolü, areas üniteyi, criteria dönemsel kriterleri, audit_plans görev ve onayı, audits tetkik sonucunu, corrective_tasks uygunsuzluk sürecini, notifications ise bildirimleri saklar.","Onaylanan planın kriterleri JSONB anlık görüntüsü olarak tutulur. Böylece sonraki değişiklikler geçmiş tetkikleri etkilemez. Kriterler alan ve YYYY-AA dönemiyle ayrıştırılır; Ağustos ayında yapılan değişiklik Eylül kriterlerini değiştirmez."],None),
('9. Kimlik Doğrulama ve Yetkilendirme',["Girişten sonra backend imzalı JWT üretir. Frontend tokenı Authorization: Bearer başlığıyla gönderir. AuthInterceptor tokenı doğrular ve kullanıcının aktifliğini veritabanından kontrol eder.","Global Admin tüm sistemi yönetir. Ünite Sorumlusu kendi ünitesinde kriter ve plan hazırlar. Tetkikçi onaylanmış planı uygular. Alan Sorumlusu yalnızca kendisine atanan uygunsuzluğu giderir. Alan sınırları yalnızca arayüzde değil, backend sorgularında da uygulanır."],('9.1',5,'Rol tabanlı yetkilendirme tablosu')),
('10. Kriter ve Planlama Modülü',["Ünite sorumlusu alan ve ay seçerek kriterleri oluşturur, ardından tetkik tarihi, tetkikçi ve alan sorumlusunu birlikte seçip planı admin onayına gönderir. Admin onayı gelmeden plan tetkikçiye açılmaz ve görevlilere bildirim gitmez.","Yanlış oluşturulan ve başlanmamış plan silinebilir. Başlamış plan, tetkik geçmişinin bütünlüğünü korumak için silinemez. Silinen ay yeniden planlanırsa yeni sürüm tekrar admin onayına gider."],('10.1',12,'Tetkik planlama ve görevli atama paneli')),
('11. QR Kodlu Tetkik ve Puanlama',["Tetkikçi atanmış planı açar veya alan QR kodunu kamerayla okutur. Her kriter Uygun ya da Uygun Değil olarak işaretlenir. Uygun olmayan maddelerde açıklama ve fotoğraf zorunludur.","Puan, uygun kriterlerin ağırlık toplamının tüm kriter ağırlıklarına bölünmesi ve 100 ile çarpılmasıyla hesaplanır. Kriterler eşit ağırlıktaysa sonuç Uygun Kriter Sayısı / Değerlendirilen Kriter Sayısı x 100 formülüyle aynıdır."],('11.1',14,'Kriter değerlendirme ve anlık puan paneli')),
('12. Düzeltici Faaliyet Akışı',["Tetkik tamamlandığında Uygun Değil kriterler için görev oluşturulur ve admin onaylı plandaki alan sorumlusuna atanır. Alan sorumlusu tetkikçinin açıklamasını ve önce fotoğrafını görür.","Sorun giderildiğinde alan sorumlusu açıklama ve sonrası fotoğrafını ekler. Görev onay bekliyor durumuna geçer; ünite sorumlusuna ve global admine bildirim gider. Yetkili kullanıcı düzeltmeyi onaylar veya gerekçeyle geri gönderir."],('12.1',16,'Alan sorumlusunun düzeltme ve kanıt paneli')),
('13. Raporlama ve Bildirim',["Raporlar tarih, alan ve durumla filtrelenebilir. Global admin tüm kayıtları, ünite sorumlusu yalnızca kendi alanını görür. Sonuçlar tarayıcının yazdırma özelliğiyle PDF olarak kaydedilebilir.","Bildirim paneli yeni planı, düzeltme görevini, son tarihi ve onayı ilgili kullanıcıya gösterir. Eğitim/Bilgilendirme ekranı 5S sunumunu, Yardım ekranı ise role özel kullanım adımlarını içerir."],('13.1',18,'Geçmiş belgelemelerin görüntülenmesi')),
('14. Test ve Güvenlik',["Frontend TypeScript kontrolü ve üretim derlemesiyle doğrulanmıştır. Backend testlerinde gerçek PostgreSQL bağlantısı üzerinde transaction rollback kullanılan geçici kayıtlar oluşturulmuştur.","Tetkikçinin görevli atayamaması, ünite dışı kullanıcı seçilememesi, admin onayından önce planın açılamaması, kriter anlık görüntüsünün korunması, başlamış planın silinememesi ve raporların alanla sınırlandırılması test edilmiştir. SQL sorguları parametreli çalışır; parola ve JWT sırrı .env değişkenlerinden okunur."],None),
('15. Sonuç',["Proje sonunda Excel ve dağınık fotoğraf kayıtlarıyla yürütülen 5S süreci rol tabanlı bir web uygulamasında birleştirilmiştir. Sistem dönemsel kriterleri, admin onaylı görevlendirmeyi, QR destekli saha tetkikini, fotoğraflı düzeltmeyi ve geçmiş raporlamayı destekler.","Çalışma sırasında React ve TypeScript ile arayüz geliştirme, Java ve Spring Boot ile REST API hazırlama, PostgreSQL veri modelleme, JWT kimlik doğrulama ve dosya yükleme konularında uygulamalı deneyim kazandım. Gelecekte kurumsal kimlik, e-posta ve ERP entegrasyonları eklenebilir."],None)
]
for name,paras,fig in sections:
 page();heading(name)
 for x in paras:para(x)
 if name.startswith('5.'):
  table(['Katman','Teknoloji','Görev'],[['Frontend','React, TypeScript, HTML, CSS','Ekran ve API çağrıları'],['Backend','Java, Spring Boot, Spring JDBC','İş kuralları ve REST servisleri'],['Veritabanı','PostgreSQL, SQL','Kalıcı ve ilişkisel kayıtlar']])
 if name.startswith('8.'):
  table(['Tablo','İçerik'],[['users','Kullanıcı ve rol'],['areas','Ünite ve yöneticiler'],['criteria','Alan ve dönem kriteri'],['audit_plans','Görevli, tarih ve admin onayı'],['audits','Tetkik sonucu'],['corrective_tasks','Uygunsuzluk ve düzeltme'],['notifications','Kullanıcı bildirimi']])
 if fig:figure(*fig)

page();heading('Ek A. Seçilmiş Kod Blokları')
code('Kod A.1 PostgreSQL bağlantısı','String value = environment.getProperty("DATABASE_URL");\nPGSimpleDataSource source = new PGSimpleDataSource();\nsource.setUrl(value);\nreturn source;')
code('Kod A.2 JWT kimlik doğrulama','String authorization = request.getHeader("Authorization");\nMap<String,Object> payload = jwt.verify(token);\nrequest.setAttribute(USER, users.get(0));')
code('Kod A.3 React ile tetkik kaydı','await fetch(`${API_BASE}/api/audits/current`, {\n  method: "POST",\n  headers: {...authHeaders, "content-type": "application/json"},\n  body: JSON.stringify({rows, score, submit, planId})\n});')
code('Kod A.4 Parametrik görev sorgusu',"SELECT t.id,t.criterion_text,t.status\nFROM corrective_tasks t\nWHERE t.assigned_to=?::uuid AND t.status='open'\nORDER BY t.due_at")

page();heading('Ek B. Panel Görselleri')
for no,slide,cap in [('B.1',10,'Ünite sorumlusu genel bakış'),('B.2',11,'Dönemsel kriter yönetimi'),('B.3',13,'Tetkikçiye atanan plan'),('B.4',15,'Uygunsuzluk kanıtı'),('B.5',17,'Düzeltme onayı')]:figure(no,slide,cap,14.5)
d.save(OUT);print(OUT)
